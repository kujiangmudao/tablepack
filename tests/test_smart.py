# -*- coding: utf-8 -*-
"""Tests for the opt-in LLM smart layer (no network — stub clients only)."""

from __future__ import annotations

from pdf_excel.smart import (
    SmartEnhancer,
    SmartResult,
    apply_smart,
    enhance_grid,
    extract_json,
    should_enhance,
)


class StubClient:
    def __init__(self, reply):
        self.reply = reply

    def complete_json(self, prompt: str):
        if isinstance(self.reply, Exception):
            raise self.reply
        return self.reply


class CountingStub(StubClient):
    def __init__(self, reply):
        super().__init__(reply)
        self.calls = 0

    def complete_json(self, prompt: str):
        self.calls += 1
        return super().complete_json(prompt)


GRID = [
    ["", "化学成分", "", "Fe2O3"],
    ["Sample", "SiO2", "Al2O3", ""],
    ["S-01", "45.2", "22.1", "8.3"],
]


def test_extract_json_variants():
    assert extract_json('{"a": 1}') == {"a": 1}
    assert extract_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert extract_json('Sure! {"a": 1} hope it helps') == {"a": 1}
    assert extract_json("no json here") == {}
    assert extract_json("[1,2,3]") == {}


def test_enhance_grid_ok():
    client = StubClient(
        {
            "header_rows": 2,
            "columns": ["Sample", "化学成分_SiO2", "化学成分_Al2O3", "Fe2O3"],
            "warnings": ["merged caption expanded"],
        }
    )
    r = enhance_grid(GRID, "表1 测试", client)
    assert r.header_rows == 2
    assert r.columns[1] == "化学成分_SiO2"
    assert "merged caption expanded" in r.warnings

    new_grid = apply_smart(GRID, r)
    assert new_grid[0][1] == "化学成分_SiO2"
    assert new_grid[1] == GRID[2]  # data rows untouched


def test_enhance_grid_bad_reply_never_raises():
    r = enhance_grid(GRID, "", StubClient({"nope": True}))
    assert r.columns == []
    assert any("columns" in w for w in r.warnings)
    assert apply_smart(GRID, r) == GRID  # unchanged


def test_enhance_grid_client_error_never_raises():
    r = enhance_grid(GRID, "", StubClient(RuntimeError("boom")))
    assert r.columns == []
    assert any("llm call failed" in w for w in r.warnings)


def test_columns_clamped_and_padded():
    client = StubClient({"header_rows": 1, "columns": ["A", "B", "C", "D", "E"]})
    r = enhance_grid(GRID, "", client)
    assert len(r.columns) == len(GRID[0])
    assert any("truncated" in w for w in r.warnings)

    client2 = StubClient({"header_rows": 1, "columns": ["A"]})
    r2 = enhance_grid(GRID, "", client2)
    assert len(r2.columns) == len(GRID[0])
    assert r2.columns[1] == ""
    assert any("padded" in w for w in r2.warnings)


def test_header_rows_clamped():
    client = StubClient({"header_rows": 99, "columns": ["a", "b", "c", "d"]})
    r = enhance_grid(GRID, "", client)
    assert r.header_rows == len(GRID)


def test_apply_smart_requires_columns():
    assert apply_smart(GRID, SmartResult()) == GRID


def test_should_enhance_detects_header_gaps():
    assert should_enhance(GRID) is True
    assert should_enhance([["a", "b"], ["1", "2"]]) is False
    assert should_enhance([["a"]]) is False


def test_smart_enhancer_skips_clean_grid():
    client = CountingStub({"header_rows": 1, "columns": ["x", "y"]})
    enh = SmartEnhancer(client)
    grid, warns = enh.enhance([["a", "b"], ["1", "2"]], "cap")
    assert client.calls == 0
    assert grid == [["a", "b"], ["1", "2"]]
    assert warns == []
    assert enh.enhanced_count == 0


def test_smart_enhancer_counts_and_warns():
    client = CountingStub({"header_rows": 2, "columns": ["a", "b", "c", "d"]})
    enh = SmartEnhancer(client)
    grid, warns = enh.enhance(GRID, "表1")
    assert client.calls == 1
    assert enh.enhanced_count == 1
    assert grid[0] == ["a", "b", "c", "d"]
    assert grid[1] == GRID[2]


def test_smart_enhancer_failure_keeps_grid():
    client = CountingStub(RuntimeError("boom"))
    enh = SmartEnhancer(client)
    grid, warns = enh.enhance(GRID, "")
    assert grid == GRID
    assert any("llm call failed" in w for w in warns)
    assert enh.enhanced_count == 0


def test_write_excel_applies_smart(tmp_path):
    from openpyxl import load_workbook

    from pdf_excel.excel_writer import write_excel
    from pdf_excel.models import TableItem

    html = (
        "<table><tr><td></td><td>化学成分</td><td></td></tr>"
        "<tr><td>S-01</td><td>45</td><td>22</td></tr></table>"
    )
    item = TableItem(
        index=1,
        page_idx=0,
        caption="表1",
        caption_raw=[],
        footnote=[],
        html_body=html,
        img_path=None,
        bbox=None,
    )
    client = CountingStub({"header_rows": 1, "columns": ["Sample", "SiO2", "Al2O3"]})
    xlsx = tmp_path / "t.xlsx"
    write_excel([item], xlsx, smart=SmartEnhancer(client))
    assert client.calls == 1
    ws = load_workbook(xlsx).active
    assert ws.cell(row=4, column=2).value == "SiO2"
    assert ws.cell(row=5, column=1).value == "S-01"


def test_make_smart_enhancer_requires_config():
    from pdf_excel.config import Settings
    from pdf_excel.pipeline import _make_smart_enhancer

    issues: list[str] = []
    assert _make_smart_enhancer(Settings(smart_enabled=True), issues) is None
    assert any("missing config" in i for i in issues)


def _mini_item():
    from pdf_excel.models import TableItem

    return TableItem(
        index=1,
        page_idx=0,
        caption="表1",
        caption_raw=[],
        footnote=[],
        html_body="<table><tr><td>a</td></tr></table>",
        img_path=None,
        bbox=None,
    )


def test_split_smart_marker():
    from pdf_excel.pipeline import _split_smart

    hard, smart = _split_smart(["a", "表2 (c): smart: x", "b", "表3: smart: y"])
    assert hard == ["a", "b"]
    assert smart == ["表2 (c): smart: x", "表3: smart: y"]


def test_write_notes_smart_notes_keep_success_file(tmp_path):
    from pdf_excel.pipeline import write_notes

    out = tmp_path / "pkg"
    out.mkdir()
    path = write_notes(
        out,
        "stem",
        "stem.pdf",
        [_mini_item()],
        0,
        [],
        smart_notes=["表1 (表1): smart: headers flattened"],
    )
    assert path.name == "转换说明.md"
    text = path.read_text(encoding="utf-8")
    assert "AI 修正备注" in text and "headers flattened" in text
    assert not (out / "问题说明.md").exists()


def test_write_notes_hard_issue_flips_and_keeps_smart(tmp_path):
    from pdf_excel.pipeline import write_notes

    out = tmp_path / "pkg"
    out.mkdir()
    path = write_notes(
        out,
        "stem",
        "stem.pdf",
        [_mini_item()],
        0,
        ["real problem"],
        smart_notes=["表1: smart: note"],
    )
    assert path.name == "问题说明.md"
    text = path.read_text(encoding="utf-8")
    assert "real problem" in text
    assert "AI 修正备注" in text and "smart: note" in text
